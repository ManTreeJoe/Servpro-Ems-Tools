"""Explicit, local run-history backfill. Does not audit jobs or write Job Logs."""
import argparse
import hashlib
import json
import os
import re
import sqlite3
from datetime import date, timedelta
from pathlib import Path


def client_key(name):
    import persistence
    value = persistence._canon_pin_key(name)
    if ',' in value:
        last, first = value.split(',', 1)
        value = first.strip() + ' ' + last.strip()
    return re.sub(r'\s+', ' ', value).strip().casefold()


def location():
    import config
    import paths
    import run_doc
    root = os.path.normcase(os.path.abspath(run_doc._runs_dir()))
    scope = hashlib.sha256((str(config.active_department()) + '\n' + root).encode()).hexdigest()
    return Path(paths.DATA_DIR) / 'run_activity_index.sqlite3', scope, root


def connect(path):
    connection = sqlite3.connect(path, timeout=3)
    connection.execute('CREATE TABLE IF NOT EXISTS run_activity '
                       '(scope TEXT, identity TEXT, client_key TEXT, day TEXT, payload TEXT, '
                       'PRIMARY KEY(scope,identity))')
    connection.execute('CREATE INDEX IF NOT EXISTS run_activity_client ON run_activity(scope,client_key,day)')
    return connection


def commit(rows, db_path, scope):
    connection = connect(db_path)
    try:
        before = connection.total_changes
        with connection:
            for row in rows:
                identity = hashlib.sha256(json.dumps([row['iso'], row['key'], row['raw']], ensure_ascii=False).encode()).hexdigest()
                connection.execute('INSERT OR IGNORE INTO run_activity VALUES (?,?,?,?,?)',
                                   (scope, identity, row['key'], row['iso'], json.dumps(row)))
        return connection.total_changes - before
    finally:
        connection.close()


def history(client):
    db_path, scope, _ = location()
    if not db_path.exists():
        return []
    connection = sqlite3.connect(f'{db_path.as_uri()}?mode=ro', uri=True, timeout=2)
    try:
        return [json.loads(row[0]) for row in connection.execute(
            'SELECT payload FROM run_activity WHERE scope=? AND client_key=? ORDER BY day DESC,identity',
            (scope, client_key(client)))]
    finally:
        connection.close()


def parse_source(path):
    """No yesterday-date fallback, merging, corpus crawl, or timestamp write."""
    import run_doc
    if path.lower().endswith('.msg'):
        import msg_reader
        entries = [(line, False) for line in msg_reader.read_msg_text(path).splitlines()]
    else:
        from docx import Document
        from audit_logic import para_is_struck
        from shared_file_read import open_read
        with open_read(path) as stream:
            document = Document(stream)
        entries = [(p.text, para_is_struck(p)) for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    entries.extend((p.text, para_is_struck(p)) for p in cell.paragraphs)
    return run_doc._parse_run_doc_entries(entries)


def scan(start, end, root):
    import run_doc
    from audit_logic import detect_activity
    from job_history import _date_from_docstr
    folders = set()
    day = start
    while day <= end:
        folders.update(run_doc._month_search_dirs(root, day.year, day.month))
        day += timedelta(days=1)
    candidates, errors, excluded, rows = {}, [], [], []
    for folder in sorted(folders):
        try:
            for name in os.listdir(folder):
                if name.startswith('~$') or not name.lower().endswith(('.docx', '.msg')):
                    continue
                day = start
                while day <= end:
                    if run_doc._run_doc_date_pattern(day).search(name):
                        candidates[os.path.join(folder, name)] = day
                        break
                    day += timedelta(days=1)
        except OSError as error:
            errors.append({'source': folder, 'error': str(error)})
    seen = set()
    for path, day in sorted(candidates.items()):
        try:
            jobs, header = parse_source(path)
            header_day = _date_from_docstr(header)
            if header_day and header_day.date() != day:
                excluded.append({'source': path, 'reason': 'Header date differs from filename', 'header': header})
                continue
            if not jobs:
                excluded.append({'source': path, 'reason': 'No active work/monitor rows parsed'})
            for job in jobs:
                if job.get('unit') or job.get('claim_hint'):
                    excluded.append({'source': path, 'client': job['client'], 'reason': 'Unit/claim needs explicit job match'})
                    continue
                key = client_key(job['client'])
                raw = job.get('raw', '').strip()
                identity = (day.isoformat(), key, raw)
                if not key or identity in seen:
                    continue
                seen.add(identity)
                info = detect_activity(raw, section=job.get('section'), new_loss=job.get('new_loss'))
                rows.append({'key': key, 'client': job['client'], 'iso': day.isoformat(),
                             'raw': raw, 'labels': info.get('labels') or [], 'techs': job.get('techs') or [],
                             'source': path, 'section': job.get('section'), 'status': 'scheduled'})
        except Exception as error:
            errors.append({'source': path, 'error': str(error)})
    return {'start': start.isoformat(), 'end': end.isoformat(), 'documents': len(candidates),
            'rows': rows, 'errors': errors, 'excluded': excluded}


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--start', required=True)
    parser.add_argument('--end', required=True)
    parser.add_argument('--apply', action='store_true')
    args = parser.parse_args()
    db_path, scope, root = location()
    report = scan(date.fromisoformat(args.start), date.fromisoformat(args.end), root)
    report['inserted'] = commit(report['rows'], db_path, scope) if args.apply else 0
    report_path = db_path.parent / f'run-history-{args.start}-{args.end}{"-applied" if args.apply else "-preview"}.json'
    report_path.write_text(json.dumps(report, indent=2), encoding='utf-8')
    print(json.dumps({'documents': report['documents'], 'entries': len(report['rows']),
                      'clients': len({r['key'] for r in report['rows']}), 'inserted': report['inserted'],
                      'errors': len(report['errors']), 'excluded': len(report['excluded']),
                      'report': str(report_path)}, indent=2))
