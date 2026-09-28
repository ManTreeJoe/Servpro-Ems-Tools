import ctypes
import os
import pytest
from docx import Document
import run_doc_editor


@pytest.mark.skipif(os.name != 'nt', reason='Windows Office sharing semantics')
def test_schedule_reads_document_held_with_delete_access(tmp_path):
    from ctypes import wintypes
    path = tmp_path / 'Wednesday.docx'
    doc = Document()
    doc.add_paragraph('Monitor')
    doc.add_paragraph('Fixture: monitor')
    doc.save(path)
    kernel = ctypes.WinDLL('kernel32', use_last_error=True)
    kernel.CreateFileW.argtypes = [wintypes.LPCWSTR,wintypes.DWORD,wintypes.DWORD,wintypes.LPVOID,wintypes.DWORD,wintypes.DWORD,wintypes.HANDLE]
    kernel.CreateFileW.restype = wintypes.HANDLE
    kernel.CloseHandle.argtypes = [wintypes.HANDLE]
    handle = kernel.CreateFileW(str(path), 0x80000000 | 0x10000, 7, None, 3, 0x80, None)
    assert handle != ctypes.c_void_p(-1).value
    try:
        with pytest.raises(PermissionError):
            with open(path, 'rb'):
                pass
        model = run_doc_editor.read_document(str(path))
        assert model['sections']['monitor'][0]['text'] == 'Fixture: monitor'
        assert model['version']
    finally:
        kernel.CloseHandle(handle)
