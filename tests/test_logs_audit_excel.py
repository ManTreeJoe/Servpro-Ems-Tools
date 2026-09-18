from copy import deepcopy
import hashlib
import pytest
from logs_audit_excel import records_for_period, save_new_copy

def draft(start='2026-09-14',end='2026-09-20',version='v1'):
    return {'version':version,'saved_at':'2026-09-17','fields':{'period_start':start,'period_end':end,'ems_estimator':'Test'},
            'evidence':{'card':{'name':'Sample','id':'card1'}},'history':[]}

def test_export_includes_moved_cards_and_period_history():
    old=draft();current=draft('2026-09-21','2026-09-27','v2');current['history']=[old]
    opid=hashlib.sha256(b'card1v1').hexdigest()[:24]
    store={'drafts':{'card1':current},'operations':{opid:{'status':'done'}}}
    result=records_for_period(store,'2026-09-14','2026-09-20')
    assert result[0]['draft']['version']=='v1'
    assert result[0]['status']=='Published'
    assert records_for_period(store,'2026-08-01','2026-08-07')==[]

def test_new_revision_is_not_mislabelled_published():
    d=draft(version='v2');opid=hashlib.sha256(b'card1v1').hexdigest()[:24]
    store={'drafts':{'card1':d},'operations':{opid:{'status':'done'}}}
    assert records_for_period(store,'2026-09-14','2026-09-20')[0]['status']=='Draft / not published'

def test_no_source_overwrite(tmp_path):
    source=tmp_path/'template.xlsx'
    with pytest.raises(ValueError,match='source template'):
        save_new_copy(source,source,[],'2026-09-14','2026-09-20','IE')

def test_export_days_reference_separate_start_inputs(tmp_path):
    import io
    from datetime import datetime
    import openpyxl
    from logs_audit_excel import HEADERS, build_copy
    wb=openpyxl.Workbook();ws=wb.active;ws.title='Weekly Audit'
    for col,label in enumerate(HEADERS,1):ws.cell(7,col,label)
    template=tmp_path/'template.xlsx';wb.save(template)
    d=draft();d['fields'].update(ems_start='2026-06-01',ems_ready='2026-06-07',ems_billed='2026-06-08')
    record={'card_id':'card1','draft':d,'status':'Draft / not published','operation':{}}
    result=openpyxl.load_workbook(io.BytesIO(build_copy(template,[record],'2026-09-14','2026-09-20','IE')))
    assert result['Weekly Audit']['J8'].value=='=IF(OR(\'Audit detail\'!C2="",I8=""),"",I8-\'Audit detail\'!C2)'
    assert result['Audit detail']['C2'].value==datetime(2026,6,1)
    assert result['Audit detail']['C3'].value is None
    assert result['Weekly Audit']['H8'].value==datetime(2026,6,7)
    assert result['Weekly Audit'].max_column==15
