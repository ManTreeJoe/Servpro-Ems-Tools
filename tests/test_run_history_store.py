from datetime import date
from docx import Document
import pytest
import run_history_store as history


def test_snapshot_keeps_waiting_crossed_out_and_tables(tmp_path):
    path=tmp_path/'Monday 7.6.26.docx'
    doc=Document()
    doc.add_paragraph('Date: 7/6/2026')
    doc.add_paragraph('Work To Be Performed')
    doc.add_paragraph('Jane Doe: Demo')
    doc.add_paragraph('TBS Mitigation')
    p=doc.add_paragraph(); p.add_run('Old job: waiting').font.strike=True
    doc.add_table(rows=1,cols=1).cell(0,0).text='Table source'
    doc.save(path)
    before=path.read_bytes()
    result=history.document(path,date(2026,7,6),'IE',[])
    assert result['run_date']=='2026-07-06'
    assert any(r['struck'] for r in result['rows'])
    assert any(r['raw_text']=='Table source' for r in result['rows'])
    assert not any('status' in r for r in result['rows'])
    assert path.read_bytes()==before
    assert history.document(path,date(2026,7,6),'IE',[])==result


def test_mismatching_header_is_not_assigned_to_wrong_day(tmp_path):
    path=tmp_path/'Monday 7.6.26.docx'
    doc=Document();doc.add_paragraph('Date: 7/7/2026');doc.save(path)
    with pytest.raises(ValueError,match='Header date'):
        history.document(path,date(2026,7,6),'IE',[])
