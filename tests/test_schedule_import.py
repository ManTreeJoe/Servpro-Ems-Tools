"""The real Schedule import preview is read-only and lossless about raw text."""
from datetime import date

from docx import Document
import pytest

import run_doc_editor as editor
from run_doc_editor_web import Api


@pytest.fixture
def source(tmp_path, monkeypatch):
    path = tmp_path / "Friday Run.docx"
    doc = Document()
    doc.add_paragraph("SERVPRO DAILY RUN")
    doc.add_paragraph("")
    for key, label, _pattern in editor.SECTION_DEFS:
        doc.add_paragraph(label)
        row = doc.add_paragraph(f"  Authored {key}:  leave spacing intact  ")
        if key == "monitor":
            row.runs[0].font.strike = True
    doc.add_paragraph("Same line")
    doc.add_paragraph("Same line")
    doc.save(path)
    monkeypatch.setattr("run_doc._find_run_doc_for_date", lambda day: str(path))
    monkeypatch.setattr("config.active_department", lambda: "IE")
    monkeypatch.setattr(Api, "_day", staticmethod(lambda offset=0: date(2026, 9, 18)))
    return path


def test_preview_preserves_all_sections_and_source_document(source):
    original, stamp = source.read_bytes(), source.stat().st_mtime_ns
    result = Api().preview_schedule_import()
    assert result["ok"] and result["preview_only"]
    assert not result["can_apply"]
    assert source.read_bytes() == original
    assert source.stat().st_mtime_ns == stamp
    assert [row["raw_text"] for row in result["paragraphs"]] == [
        p.text for p in Document(source).paragraphs]
    assert {v["section"] for v in result["visits"]} == set(editor.SECTIONS)
    monitor = next(v for v in result["visits"] if v["section"] == "monitor")
    assert monitor["struck"]
    assert monitor["raw_text"] == "  Authored monitor:  leave spacing intact  "
    assert monitor["proposed_date"] == "2026-09-18"
    for visit in result["visits"]:
        assert visit["review_required"]
        assert visit["job_id"] is None
        assert visit["division"] is None
        if visit["section"] not in {"work", "monitor"}:
            assert visit["proposed_date"] is None


def test_preview_keys_are_repeatable_scoped_and_duplicate_safe(source, monkeypatch):
    first, repeated = Api().preview_schedule_import(), Api().preview_schedule_import()
    assert first == repeated
    keys = [v["source_row_key"] for v in first["visits"]]
    assert len(keys) == len(set(keys))
    monkeypatch.setattr("config.active_department", lambda: "OC")
    other = Api().preview_schedule_import()
    assert other["import_key"] != first["import_key"]
    assert not set(keys).intersection(v["source_row_key"] for v in other["visits"])


def test_changed_source_requires_a_new_preview_not_position_based_updates(source):
    first = Api().preview_schedule_import()
    doc = Document(source)
    doc.paragraphs[-1].text = "Edited outside Hub"
    doc.save(source)
    changed = Api().preview_schedule_import()
    assert first["source_version"] != changed["source_version"]
    assert first["import_key"] != changed["import_key"]
    assert changed["requires_reconciliation"]


def test_tables_are_explicitly_flagged_not_silently_dropped(source):
    doc = Document(source)
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Extra dispatch instructions"
    doc.save(source)
    result = Api().preview_schedule_import()
    assert "tables_require_review" in result["blockers"]
    assert result["tables"][0][0][0] == "Extra dispatch instructions"
    assert not result["can_apply"]


def test_missing_day_does_not_create_anything(monkeypatch):
    monkeypatch.setattr("run_doc._find_run_doc_for_date", lambda day: None)
    assert not Api().preview_schedule_import()["ok"]


def test_preview_rejects_an_outdated_day_or_workspace_before_reading(source, monkeypatch):
    reads = []
    def forbidden(day):
        reads.append(day)
        return str(source)
    monkeypatch.setattr("run_doc._find_run_doc_for_date", forbidden)
    assert not Api().preview_schedule_import(0, "2026-09-17", "IE")["ok"]
    assert not Api().preview_schedule_import(0, "2026-09-18", "OC")["ok"]
    assert reads == []


def test_preview_does_not_return_old_workspace_data_after_switch(source, monkeypatch):
    import schedule_import
    original = schedule_import.preview
    def switched(*args, **kwargs):
        result = original(*args, **kwargs)
        monkeypatch.setattr("config.active_department", lambda: "OC")
        return result
    monkeypatch.setattr(schedule_import, "preview", switched)
    result = Api().preview_schedule_import(0, "2026-09-18", "IE")
    assert not result["ok"]
    assert "visits" not in result


def test_embedded_home_bridge_exposes_the_real_preview(source):
    from home_web import HomeApi
    home = object.__new__(HomeApi)
    home._bind_methods("run_doc_editor", Api())
    result = home.run_doc_editor_preview_schedule_import(0, "2026-09-18", "IE")
    assert result["ok"] and result["preview_only"]
    assert len(result["visits"]) == len(editor.SECTIONS) + 2
