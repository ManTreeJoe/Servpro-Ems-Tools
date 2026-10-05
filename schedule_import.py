"""Read-only run-document preview for the shared Schedule cutover.

Raw paragraphs remain authoritative for import review. These are proposals,
not persisted visits: no name matching, assumed crew, or implicit completion.
Source-row keys identify ONE document revision, not visits across revisions.
The later confirmed import must reconcile changed documents and assign durable
visit IDs; blindly upserting a new preview would duplicate carried-over work.
"""
from __future__ import annotations

from datetime import date
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
from zipfile import ZipFile
from uuid import NAMESPACE_URL, uuid5

from docx import Document

import run_doc_editor as editor


def preview(path: str, *, day: date, workspace: str) -> dict:
    """Read one immutable byte snapshot. Never write to Word or any database."""
    workspace = str(workspace or "").strip().upper()
    if not workspace:
        raise ValueError("A workspace is required for Schedule import")
    source = Path(path)
    if source.suffix.casefold() != ".docx":
        raise ValueError("Schedule import preview requires a Word .docx document")
    with source.open('rb') as stream:
        content = stream.read(20 * 1024 * 1024 + 1)
    if len(content) > 20 * 1024 * 1024:
        raise ValueError('Run document must be 20 MB or smaller')
    with ZipFile(BytesIO(content)) as archive:
        if sum(info.file_size for info in archive.infolist()) > 80 * 1024 * 1024:
            raise ValueError('Expanded Run document is too large')
    version = sha256(content).hexdigest()
    # Path is informational only: mapped drives and UNC paths may identify
    # the same file. Workspace/day/content are the replay identity.
    import_key = sha256(json.dumps(
        ["run-import-v1", workspace, day.isoformat(), version],
        separators=(",", ":")).encode()).hexdigest()
    document = Document(BytesIO(content))
    paragraphs, visits = [], []
    section = None
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text or ""
        # Unlike the legacy editor scanner, recognize repeated headings too.
        heading = editor._section_for(text, None)
        if heading:
            section = heading
        kind = ("heading" if heading else "blank" if not text.strip()
                else "candidate" if section else "preserved")
        raw = {"index": index, "kind": kind, "section": section,
               "raw_text": text, "struck": editor._paragraph_is_struck(paragraph)}
        paragraphs.append(raw)
        if kind != "candidate":
            continue
        visits.append({
            "source_row_key": f"{import_key}:{index}",
            "source_paragraph": index, "section": section,
            "position": len(visits), "raw_text": text, "struck": raw["struck"],
            "proposed_date": day.isoformat() if section in {"work", "monitor"} else None,
            "job_id": None, "division": None, "crew": [],
            "arrival_start": None, "arrival_end": None,
            "review_required": True,
        })
    tables = [[[cell.text for cell in row.cells] for row in table.rows]
              for table in document.tables]
    blockers = ["shared_schedule_cutover_pending", "confirm_job_links_and_visit_details"]
    if tables:
        blockers.append("tables_require_review")
    if document.element.xpath(".//w:ins | .//w:del | .//w:txbxContent"):
        blockers.append("tracked_changes_or_textboxes_require_review")
    if not visits:
        blockers.append("no_recognized_run_rows")
    return {
        "ok": True, "preview_only": True, "can_apply": False,
        "requires_reconciliation": True, "workspace": workspace,
        "run_date": day.isoformat(), "source_path": str(source.absolute()),
        "source_filename": source.name, "section_labels": dict(editor.SECTION_LABELS),
        "source_version": version, "source_bytes": len(content),
        "import_key": import_key, "paragraphs": paragraphs,
        "visits": visits, "tables": tables, "blockers": blockers,
        "print_template_required": True,
    }


def draft_rows(result):
    """Proposals only: explicit job selection and normal draft save are required."""
    mapping = {
        'work': ('scheduled', 'Work To Be Performed'),
        'monitor': ('scheduled', 'Monitor'),
        'upcoming': ('scheduled', 'Work To Be Performed'),
        'tbs_new_loss': ('tbs', 'TBS New Loss /Reinspection'),
        'tbs_mitigation': ('tbs', 'TBS Mitigation'),
        'tbs_contents': ('tbs', 'TBS Contents'),
        'pending_testing': ('pending', 'Pending Testing/Clearance/Abatement'),
        'pending_insurance': ('pending', 'Pending Approvals – Insurance/Self Pay'),
        'pending_property': ('pending', 'Pending Approvals – Property Management'),
        'on_hold': ('hold', 'On Hold'),
    }
    proposals = []
    for row in result['visits']:
        placement = mapping.get(row['section'])
        reason = ('Crossed out in document' if row['struck'] else
                  'Section needs manual handling' if not placement else
                  'Line exceeds the 3,800-character import limit' if len(row['raw_text']) > 3800 else '')
        proposals.append({**row, 'queue': placement[0] if placement else None,
            'group': placement[1] if placement else None, 'skip_reason': reason,
            'draft_id': str(uuid5(NAMESPACE_URL, row['source_row_key']))})
    return {**result, 'visits': proposals}
