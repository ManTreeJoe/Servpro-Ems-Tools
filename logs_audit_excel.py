"""Installed-app weekly-copy exporter. Never writes into the source template."""
from copy import copy, deepcopy
from datetime import date
from pathlib import Path
import hashlib
import io

HEADERS = ['Job / Customer','Job Date','EMS Estimator','Contents Estimator',
           'Initial Note Sent On Time?','Coordinator Who Missed It','WC / File Status',
           'Ready for Billing','Billing Completed','EMS Days','Ready for Billing',
           'Billing Completed','Contents Days','Estimator Notes','Audit Notes / Open Action Needed']

def records_for_period(store, start, end):
    start=date.fromisoformat(start).isoformat();end=date.fromisoformat(end).isoformat()
    if start>end: raise ValueError('Invalid review period.')
    records=[]
    for card_id,current in store.get('drafts',{}).items():
        versions=[*current.get('history',[]),current]
        matches=[v for v in versions if v['fields'].get('period_start')==start and v['fields'].get('period_end')==end]
        if not matches:continue
        draft=deepcopy(matches[-1]);f=draft['fields']
        op_id=hashlib.sha256((card_id+draft['version']).encode()).hexdigest()[:24]
        operation=store.get('operations',{}).get(op_id,{})
        status='Published' if operation.get('status')=='done' else 'Draft / not published'
        records.append({'card_id':card_id,'draft':draft,'status':status,'operation':deepcopy(operation)})
    return sorted(records,key=lambda r:r['draft']['evidence']['card']['name'].casefold())

def text(cell,value):
    # External notes are literal text, never executable Excel formulas.
    cell.value=str(value or '')[:32767];cell.data_type='s'

def build_copy(template, records, start, end, workspace):
    import openpyxl
    from openpyxl.styles import Alignment, Font
    from openpyxl.workbook.properties import CalcProperties
    if not records:raise ValueError('No saved audit forms for this period. Save the job reviews first.')
    wb=openpyxl.load_workbook(template)
    if wb.sheetnames!=['Weekly Audit']:
        raise ValueError('Select the original single-sheet Weekly Audit template.')
    sheet=wb['Weekly Audit']
    if [sheet.cell(7,c).value for c in range(1,16)]!=HEADERS:
        raise ValueError('Template columns do not match the weekly review form. Nothing was exported.')
    # All rows below the template header are record slots, not prior-week data.
    last=max(sheet.max_row,7+len(records))
    styles=[copy(sheet.cell(8,c)._style) for c in range(1,16)]
    for row in sheet.iter_rows(min_row=8,max_row=last,max_col=15):
        for cell in row:cell.value=None;cell.hyperlink=None;cell.comment=None
    sheet['B2']=date.today();sheet['B2'].number_format='mm/dd/yy'
    text(sheet['E2'],f'{start} through {end}')
    text(sheet['I2'],'See Audit detail')
    detail=wb.create_sheet('Audit detail')
    detail.append(['Job / Customer','Field','Recorded value / source'])
    for cell in detail[1]:cell.font=Font(bold=True)
    detail.column_dimensions['A'].width=38;detail.column_dimensions['B'].width=25;detail.column_dimensions['C'].width=105
    detail.freeze_panes='C2'
    detail.page_setup.orientation='landscape';detail.page_setup.paperSize=sheet.page_setup.paperSize
    detail.page_setup.fitToWidth=1;detail.page_setup.fitToHeight=0
    detail.sheet_properties.pageSetUpPr.fitToPage=True
    detail.print_title_rows='1:1'
    for rowno,record in enumerate(records,8):
        d=record['draft'];f=d['fields'];card=d['evidence']['card'];name=card['name']
        for c in range(1,16):sheet.cell(rowno,c)._style=copy(styles[c-1])
        values={1:name,3:f.get('ems_estimator'),4:f.get('contents_estimator'),
                5:f.get('initial_note') or 'Unverified',6:f.get('missed_by') if f.get('initial_note')=='No' else '',
                7:f.get('file_status')}
        for col,value in values.items():text(sheet.cell(rowno,col),value)
        for col,key in ((2,'job_date'),(8,'ems_ready'),(9,'ems_billed'),(11,'contents_ready'),(12,'contents_billed')):
            sheet.cell(rowno,col).value=date.fromisoformat(f[key]) if f.get(key) else None
            sheet.cell(rowno,col).number_format='mm/dd/yy'
        # Keep the 15-column template; start dates are typed inputs in the detail tab.
        # Never substitute ready-for-billing for missing work/card start evidence.
        for prefix,col,billed_col in (('ems',10,'I'),('contents',13,'L')):
            nr=detail.max_row+1
            text(detail.cell(nr,1),name)
            text(detail.cell(nr,2),'EMS work started' if prefix=='ems' else 'Original Contents card started')
            start=f.get(prefix+'_start')
            detail.cell(nr,3).value=date.fromisoformat(start) if start else None
            detail.cell(nr,3).number_format='mm/dd/yy'
            ref=f"'Audit detail'!C{nr}"
            sheet.cell(rowno,col,f'=IF(OR({ref}="",{billed_col}{rowno}=""),"",{billed_col}{rowno}-{ref})')
        brief=lambda value: str(value or '') if len(str(value or ''))<=180 else str(value)[:150]+'… [Full text: Audit detail]'
        text(sheet.cell(rowno,14),brief(f.get('estimator_notes')))
        text(sheet.cell(rowno,15),brief(f"{record['status']}. {f.get('decision','')}\n{f.get('audit_notes','')}"))
        sheet.row_dimensions[rowno].height=110
        for cell in sheet[rowno][:15]:cell.alignment=Alignment(wrap_text=True,vertical='top')
        info={'Workspace':workspace,'Review period':f'{start} through {end}',
              'Publication status':record['status'],'Saved at':d.get('saved_at',''),
              'Logs card':card.get('shortUrl',''),'Card ID':record['card_id'],
              'AR card':(d.get('ar') or {}).get('card',{}).get('shortUrl',''),
              **f,'Published audit comment':record['operation'].get('comment','')}
        for field,suggestion in d.get('suggestions',{}).items():
            for source in suggestion.get('sources',[]):
                info[f"Source: {field} / {source['id']}"]=source.get('title','')+'\n'+source.get('text','')
        for key,value in info.items():
            if value is None or value=='':continue
            # Split long evidence into readable rows; do not truncate it.
            value=str(value)
            for offset in range(0,len(value),350):
                nr=detail.max_row+1
                for col,v in enumerate((name,key if offset==0 else key+' (continued)',value[offset:offset+350]),1):
                    text(detail.cell(nr,col),v);detail.cell(nr,col).alignment=Alignment(wrap_text=True,vertical='top')
                detail.row_dimensions[nr].height=100
    sheet.print_area=f'A1:O{7+len(records)}'
    detail.print_area=f'A1:C{detail.max_row}'
    wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
    output=io.BytesIO();wb.save(output);return output.getvalue()

def save_new_copy(template,destination,records,start,end,workspace):
    template=Path(template).resolve();destination=Path(destination).resolve()
    if destination==template:raise ValueError('Choose a new file, not the source template.')
    if destination.suffix.lower()!='.xlsx':raise ValueError('Save the copy as .xlsx.')
    payload=build_copy(template,records,start,end,workspace)
    # Refuse to overwrite any prior weekly file, including a race after the dialog.
    with destination.open('xb') as output:output.write(payload)
    return len(records)
